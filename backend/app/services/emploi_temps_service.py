"""
Service pour la gestion des emplois du temps
Fonctions de génération PDF, Excel et vérification de cohérence
"""
from datetime import date, timedelta
from typing import Optional
from io import BytesIO

from sqlalchemy.orm import Session
from sqlalchemy import func

from app.models.seance import Seance
from app.models.emploi_temps import EmploiTemps
from app.models.creneau_horaire import CreneauHoraire
from app.repositories.emploi_temps_repository import emploi_temps_repository
from app.utils.emploi_temps_utils import get_jour_semaine_label, JOURS_SEMAINE_LABELS


def generer_emploi_temps_pdf(db: Session, emploi_temps_id: int) -> bytes:
    """
    Génère un PDF de l'emploi du temps.
    
    Args:
        db: Session de base de données
        emploi_temps_id: ID de l'emploi du temps
    
    Returns:
        Contenu du PDF en bytes
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.enums import TA_CENTER
    
    # Récupérer l'emploi du temps avec les séances
    data = emploi_temps_repository.get_with_seances(db, emploi_temps_id)
    if not data:
        raise ValueError("Emploi du temps non trouvé")
    
    # Récupérer les créneaux
    creneaux = db.query(CreneauHoraire).filter(
        CreneauHoraire.is_active == True
    ).order_by(CreneauHoraire.ordre).all()
    
    # Créer le buffer PDF
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=1*cm,
        leftMargin=1*cm,
        topMargin=1*cm,
        bottomMargin=1*cm
    )
    
    elements = []
    styles = getSampleStyleSheet()
    
    # Titre
    title_style = ParagraphStyle(
        'Title',
        parent=styles['Heading1'],
        alignment=TA_CENTER,
        fontSize=16,
        spaceAfter=20
    )
    elements.append(Paragraph(f"Emploi du Temps - {data['libelle']}", title_style))
    
    # Sous-titre avec période
    subtitle_style = ParagraphStyle(
        'Subtitle',
        parent=styles['Normal'],
        alignment=TA_CENTER,
        fontSize=10,
        spaceAfter=20
    )
    elements.append(Paragraph(
        f"Semestre {data['semestre']} - Du {data['date_debut']} au {data['date_fin']}",
        subtitle_style
    ))
    elements.append(Spacer(1, 0.5*cm))
    
    # Construire le tableau
    # En-tête: Créneau | Lundi | Mardi | Mercredi | Jeudi | Vendredi | Samedi
    header = ["Créneau"] + [JOURS_SEMAINE_LABELS[i] for i in range(1, 7)]
    table_data = [header]
    
    # Organiser les séances par créneau et jour
    seances_par_creneau_jour = {}
    for seance in data.get("seances", []):
        key = (seance["creneau_id"], seance["jour_semaine"])
        if key not in seances_par_creneau_jour:
            seances_par_creneau_jour[key] = []
        seances_par_creneau_jour[key].append(seance)
    
    # Remplir le tableau
    for creneau in creneaux:
        row = [f"{creneau.libelle}"]
        for jour in range(1, 7):  # Lundi à Samedi
            key = (creneau.id, jour)
            if key in seances_par_creneau_jour:
                seances = seances_par_creneau_jour[key]
                cell_content = []
                for s in seances:
                    matiere = s.get("matiere_code", "")
                    salle = s.get("salle_code", "")
                    type_s = s.get("type_seance", "")[:2].upper()
                    cell_content.append(f"{matiere}\n{type_s} - {salle}")
                row.append("\n".join(cell_content))
            else:
                row.append("")
        table_data.append(row)
    
    # Créer le tableau
    col_widths = [3*cm] + [4*cm] * 6
    table = Table(table_data, colWidths=col_widths)
    
    # Style du tableau
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1976D2')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (0, -1), colors.HexColor('#E3F2FD')),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('ROWBACKGROUNDS', (1, 1), (-1, -1), [colors.white, colors.HexColor('#F5F5F5')]),
    ]))
    
    elements.append(table)
    
    # Légende
    elements.append(Spacer(1, 1*cm))
    legend_style = ParagraphStyle('Legend', parent=styles['Normal'], fontSize=8)
    elements.append(Paragraph("Légende: CO=Cours, TD=Travaux Dirigés, TP=Travaux Pratiques", legend_style))
    
    # Générer le PDF
    doc.build(elements)
    
    buffer.seek(0)
    return buffer.getvalue()


def generer_emploi_temps_excel(db: Session, emploi_temps_id: int) -> bytes:
    """
    Génère un fichier Excel de l'emploi du temps.
    
    Args:
        db: Session de base de données
        emploi_temps_id: ID de l'emploi du temps
    
    Returns:
        Contenu du fichier Excel en bytes
    """
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    
    # Récupérer l'emploi du temps avec les séances
    data = emploi_temps_repository.get_with_seances(db, emploi_temps_id)
    if not data:
        raise ValueError("Emploi du temps non trouvé")
    
    # Récupérer les créneaux
    creneaux = db.query(CreneauHoraire).filter(
        CreneauHoraire.is_active == True
    ).order_by(CreneauHoraire.ordre).all()
    
    # Créer le workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Emploi du Temps"
    
    # Styles
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1976D2", end_color="1976D2", fill_type="solid")
    creneau_fill = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Titre
    ws.merge_cells('A1:G1')
    ws['A1'] = f"Emploi du Temps - {data['libelle']}"
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = center_align
    
    # Sous-titre
    ws.merge_cells('A2:G2')
    ws['A2'] = f"Semestre {data['semestre']} - Du {data['date_debut']} au {data['date_fin']}"
    ws['A2'].alignment = center_align
    
    # En-tête
    headers = ["Créneau", "Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
    
    # Organiser les séances par créneau et jour
    seances_par_creneau_jour = {}
    for seance in data.get("seances", []):
        key = (seance["creneau_id"], seance["jour_semaine"])
        if key not in seances_par_creneau_jour:
            seances_par_creneau_jour[key] = []
        seances_par_creneau_jour[key].append(seance)
    
    # Remplir les données
    row_num = 5
    for creneau in creneaux:
        # Créneau
        cell = ws.cell(row=row_num, column=1, value=creneau.libelle)
        cell.fill = creneau_fill
        cell.alignment = center_align
        cell.border = thin_border
        
        # Jours
        for jour in range(1, 7):
            key = (creneau.id, jour)
            cell = ws.cell(row=row_num, column=jour + 1)
            
            if key in seances_par_creneau_jour:
                seances = seances_par_creneau_jour[key]
                content = []
                for s in seances:
                    matiere = s.get("matiere_code", "")
                    salle = s.get("salle_code", "")
                    type_s = s.get("type_seance", "")
                    content.append(f"{matiere} ({type_s})\n{salle}")
                cell.value = "\n".join(content)
            
            cell.alignment = center_align
            cell.border = thin_border
        
        row_num += 1
    
    # Ajuster la largeur des colonnes
    ws.column_dimensions['A'].width = 15
    for col in range(2, 8):
        ws.column_dimensions[get_column_letter(col)].width = 18
    
    # Ajuster la hauteur des lignes
    for row in range(5, row_num):
        ws.row_dimensions[row].height = 50
    
    # Sauvegarder dans un buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    return buffer.getvalue()


def verifier_coherence_emploi_temps(
    db: Session,
    niveau_id: int,
    filiere_id: Optional[int],
    semestre: int,
    annee_id: int
) -> dict:
    """
    Vérifie la cohérence d'un emploi du temps.
    
    Args:
        db: Session de base de données
        niveau_id: ID du niveau
        filiere_id: ID de la filière (optionnel)
        semestre: Numéro du semestre
        annee_id: ID de l'année académique
    
    Returns:
        Rapport avec erreurs et avertissements
    """
    rapport = {
        "erreurs": [],
        "avertissements": [],
        "statistiques": {}
    }
    
    # Récupérer les séances
    query = db.query(Seance).filter(
        Seance.niveau_id == niveau_id,
        Seance.semestre == semestre,
        Seance.annee_academique_id == annee_id,
        Seance.statut.notin_(["annulee"])
    )
    
    if filiere_id:
        query = query.filter(
            (Seance.filiere_id == filiere_id) | (Seance.filiere_id.is_(None))
        )
    
    seances = query.all()
    
    # Vérifier les conflits de salle
    conflits_salle = {}
    for seance in seances:
        if seance.salle_id:
            key = (seance.date_seance, seance.creneau_id, seance.salle_id)
            if key not in conflits_salle:
                conflits_salle[key] = []
            conflits_salle[key].append(seance)
    
    for key, seances_conflit in conflits_salle.items():
        if len(seances_conflit) > 1:
            rapport["erreurs"].append({
                "type": "conflit_salle",
                "message": f"Conflit de salle le {key[0]} sur le créneau {key[1]}",
                "seances": [s.code for s in seances_conflit]
            })
    
    # Vérifier les conflits d'enseignant
    conflits_enseignant = {}
    for seance in seances:
        key = (seance.date_seance, seance.creneau_id, seance.enseignant_id)
        if key not in conflits_enseignant:
            conflits_enseignant[key] = []
        conflits_enseignant[key].append(seance)
    
    for key, seances_conflit in conflits_enseignant.items():
        if len(seances_conflit) > 1:
            rapport["erreurs"].append({
                "type": "conflit_enseignant",
                "message": f"Conflit d'enseignant le {key[0]} sur le créneau {key[1]}",
                "seances": [s.code for s in seances_conflit]
            })
    
    # Vérifier la surcharge horaire (plus de 8h par jour)
    seances_par_jour = {}
    for seance in seances:
        if seance.date_seance not in seances_par_jour:
            seances_par_jour[seance.date_seance] = 0
        seances_par_jour[seance.date_seance] += seance.duree_minutes
    
    for jour, duree in seances_par_jour.items():
        if duree > 480:  # 8 heures
            rapport["avertissements"].append({
                "type": "surcharge_horaire",
                "message": f"Surcharge horaire le {jour}: {duree // 60}h{duree % 60:02d}"
            })
    
    # Statistiques
    rapport["statistiques"] = {
        "total_seances": len(seances),
        "total_heures": sum(s.duree_minutes for s in seances) // 60,
        "seances_par_type": {},
        "jours_avec_cours": len(seances_par_jour)
    }
    
    for seance in seances:
        type_s = seance.type_seance
        if type_s not in rapport["statistiques"]["seances_par_type"]:
            rapport["statistiques"]["seances_par_type"][type_s] = 0
        rapport["statistiques"]["seances_par_type"][type_s] += 1
    
    return rapport


def calculer_charge_enseignant(
    db: Session,
    enseignant_id: int,
    annee_id: int
) -> dict:
    """
    Calcule la charge horaire d'un enseignant.
    
    Args:
        db: Session de base de données
        enseignant_id: ID de l'enseignant
        annee_id: ID de l'année académique
    
    Returns:
        Statistiques de charge horaire
    """
    seances = db.query(Seance).filter(
        Seance.enseignant_id == enseignant_id,
        Seance.annee_academique_id == annee_id,
        Seance.statut.notin_(["annulee"])
    ).all()
    
    if not seances:
        return {
            "heures_totales": 0,
            "heures_semaine_moyenne": 0,
            "nombre_seances": 0,
            "repartition_par_type": {},
            "repartition_par_niveau": {}
        }
    
    # Calculer les heures totales
    heures_totales = sum(s.duree_minutes for s in seances) // 60
    
    # Calculer le nombre de semaines
    dates = [s.date_seance for s in seances]
    if dates:
        date_min = min(dates)
        date_max = max(dates)
        nb_semaines = max(1, (date_max - date_min).days // 7 + 1)
    else:
        nb_semaines = 1
    
    heures_semaine = heures_totales / nb_semaines
    
    # Répartition par type
    repartition_type = {}
    for seance in seances:
        type_s = seance.type_seance
        if type_s not in repartition_type:
            repartition_type[type_s] = {"nombre": 0, "heures": 0}
        repartition_type[type_s]["nombre"] += 1
        repartition_type[type_s]["heures"] += seance.duree_minutes // 60
    
    # Répartition par niveau
    repartition_niveau = {}
    for seance in seances:
        niveau_id = seance.niveau_id
        if niveau_id not in repartition_niveau:
            repartition_niveau[niveau_id] = {"nombre": 0, "heures": 0}
        repartition_niveau[niveau_id]["nombre"] += 1
        repartition_niveau[niveau_id]["heures"] += seance.duree_minutes // 60
    
    return {
        "heures_totales": heures_totales,
        "heures_semaine_moyenne": round(heures_semaine, 1),
        "nombre_seances": len(seances),
        "nombre_semaines": nb_semaines,
        "repartition_par_type": repartition_type,
        "repartition_par_niveau": repartition_niveau
    }
